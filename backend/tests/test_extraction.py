from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from app.services.ingestion.extraction import extract_text


def test_extract_txt():
    content = (
        b"Customer support is available Monday to Friday."
    )

    result = extract_text(
        content,
        "support.txt",
    )

    assert result == (
        "Customer support is available Monday to Friday."
    )


def test_extract_markdown():
    content = b"# Customer Support\n\nSupport is available 8 AM to 5 PM."

    result = extract_text(
        content,
        "support.md",
    )

    assert result == (
        "# Customer Support\n\n"
        "Support is available 8 AM to 5 PM."
    )


def test_extract_docx():
    document = Document()

    document.add_paragraph("Customer Support FAQ")
    document.add_paragraph(
        "Customer support is available Monday to Friday."
    )

    buffer = BytesIO()
    document.save(buffer)

    result = extract_text(
        buffer.getvalue(),
        "support.docx",
    )

    assert result == (
        "Customer Support FAQ\n\n"
        "Customer support is available Monday to Friday."
    )


def test_extract_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)

    buffer = BytesIO()
    writer.write(buffer)

    result = extract_text(
        buffer.getvalue(),
        "support.pdf",
    )

    assert result == ""


def test_extract_rejects_unsupported_file_type():
    with pytest.raises(
        ValueError,
        match="Unsupported file type",
    ):
        extract_text(
            b"some content",
            "support.xlsx",
        )
